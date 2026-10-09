"""Investigation only: export explicitly selected worksheet regions for profiling.

Run with uv run --no-project --with openpyxl python export_workbooks.py
    --manifest CAPTURES --selections SELECTIONS --out DIRECTORY.
This is not a platform reader. Source cell values and original coordinates are
retained; no names, identifiers, rows, dates or types are normalized.
"""
import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def export(manifest, selections, out):
    out.mkdir(parents=True, exist_ok=False)
    result = []
    for capture in json.loads(manifest.read_text()):
        path = Path(capture['local_file'])
        assert capture['http_status'] == 200 and sha256(path) == capture['sha256']
        workbook = load_workbook(path, read_only=True, data_only=False)
        assert set(workbook.sheetnames) == set(selections[path.name])
        for sheet in workbook:
            header_row = selections[path.name][sheet.title]
            grid = list(sheet.iter_rows(values_only=True))
            assert not any(cell.data_type in {'f', 'e'} for row in sheet for cell in row), 'formula/error needs review'
            if header_row:
                headers = list(grid[header_row - 1])
                # Blank source header cells get coordinate names, never a guessed meaning.
                columns = [str(v) if v is not None else f'unnamed_column_{i + 1}' for i, v in enumerate(headers)]
                rows = grid[header_row:]
            else:
                columns = [f'column_{i + 1}' for i in range(sheet.max_column)]
                rows = grid
            assert len(set(columns)) == len(columns), 'duplicate header needs review'
            target = out / f'{path.stem}-{workbook.sheetnames.index(sheet.title) + 1}.jsonl'
            objects = [dict(zip(columns, row, strict=True)) for row in rows]
            target.write_text(''.join(json.dumps(row, ensure_ascii=False) + '\n' for row in objects))
            roundtrip = [json.loads(line) for line in target.read_text().splitlines()]
            assert [tuple(row[col] for col in columns) for row in roundtrip] == rows
            # Entire original grid is pinned beside the projection, including preambles.
            audit = target.with_suffix('.grid.json')
            audit.write_text(json.dumps({'sheet': sheet.title, 'rows': grid}, ensure_ascii=False))
            result.append({'file': str(target), 'sha256': sha256(target), 'bytes': target.stat().st_size,
                           'raw_sha256': capture['sha256'], 'raw_file': path.name,
                           'sheet': sheet.title, 'sheet_rows': len(grid), 'sheet_columns': sheet.max_column,
                           'header_row': header_row, 'first_source_row': header_row + 1,
                           'projected_rows': len(rows), 'blank_rows_retained': sum(all(v is None for v in row) for row in rows),
                           'columns': columns, 'grid_sha256': sha256(audit), 'roundtrip_cells_equal': True})
        workbook.close()
    (out / 'derived-manifest.json').write_text(json.dumps(result, indent=2) + '\n')
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--manifest', type=Path, required=True)
    parser.add_argument('--selections', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(export(args.manifest, json.loads(args.selections.read_text()), args.out), indent=2))
