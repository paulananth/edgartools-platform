"""The Mapping Document: a spreadsheet generated from the rules (rules skill
ticket 11).

One workbook per source (`rules/sources/<source>/MAPPING.xlsx`) and one per
kind (`rules/merge/kinds/<kind>.xlsx`). Every sheet but Notes is generated
from the rules; stewards change it, `diff` prints what they changed, cell by
cell, and Claude turns that into the rules. The Notes sheet is the stewards'
own and is kept as written when the workbook is regenerated. `check` fails
when a workbook differs from what its rules generate, so the two cannot
drift.

Operator, 2026-09-28: "Initially, machine-generated must be able to edit and
update by stewards in an easily understandable way for humans".
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from . import files

NOTES = "Notes"
NOTES_HEADER = ["Note", "About", "Who", "When"]
MAX_LISTED = 12  # longer lists show as a count, and the rules file holds them
JOINING = {"cik", "lei"}  # the only identifiers that join two records into one

HOW_TO_CHANGE = (
    "Generated from the rules. To change a rule, change its cell here, then run "
    "`edgar-warehouse rules mapdoc diff`; Claude turns the change into the rules, "
    "and it acts only once its rules version is approved. Write your own notes on "
    "the Notes sheet: they are kept."
)

# What each data quality fix and test does, in plain words.
QUALITY_WORDS = {
    "blank_values@1": "Empties a value the source writes wrongly: {values}",
    "name_state_marker@1": "Fills an empty {target} from the US state tag at the end of the name (\"/DE\")",
    "standardize_address@1": "A copy for matching only: upper case, USPS street words, no suite, 5-digit ZIP",
    "present@1": "The value is filled",
    "in_set@1": "The value is one of: {values}",
    "pattern@1": "The value matches {regex}",
    "lei_check_digit@1": "The LEI check digits pass",
    "placeholder@1": "The value is not a placeholder: {values}",
    "registered_agent_address@1": "The address is not a registered agent's: {markers}",
    "in_reference@1": "The value is in the reference table {table}",
}
ON_FAIL_WORDS = {
    "exception": "Exception: the record never merges and never stops the run; fix or ignore it",
    "withhold": "Withheld: kept on the record, never used to match",
    "flag": "Counted only",
}


def _text(value: Any) -> str:
    """One cell's text: a list longer than MAX_LISTED shows as a count."""
    if value is None:
        return ""
    if isinstance(value, list):
        if len(value) > MAX_LISTED:
            return f"{len(value)} values (see the rules file)"
        return ", ".join(_text(v) for v in value)
    if isinstance(value, dict):
        return "; ".join(f"{k}: {_text(v)}" for k, v in value.items())
    return str(value)


def _words(table: dict, name: str, args: dict) -> str:
    template = table.get(name)
    if template is None:
        return f"{name} {_text(args)}".strip()
    return template.format(**{k: _text(v) for k, v in args.items()}) if args else template.split(":")[0]


def _paths(spec: Any) -> list[tuple[str, str]]:
    """(part, source path) for one mapped value: a path, or an address's parts."""
    if isinstance(spec, str):
        return [("", spec)]
    parts = []
    for part, path in (spec.get("components") or {}).items():
        parts.append((part, path["lines"] + " (lines)" if isinstance(path, dict) else path))
    return parts


# --- a source's workbook ----------------------------------------------------------

def _source_sheets(name: str, body: dict, policy: dict) -> dict[str, list[list[str]]]:
    datasets = (body.get("mdm") or {}).items()
    about = [["What", "Value"], ["How to change this", HOW_TO_CHANGE], ["Source", name],
             ["Captured files (bronze family)", _text((body.get("bronze") or {}).get("family"))]]
    fields = [["Dataset", "MDM field", "Part", "Source path", "Used for", "Rules path"]]
    identifiers = [["Dataset", "Identifier", "Source path", "Format", "Joins records into one",
                    "Can it change", "Rules path"]]
    critical = [["Dataset", "MDM field", "When it is missing", "Rules path"]]
    quality = [["Dataset", "Id", "Fix or check", "What it does", "Reads", "When it fails", "Rules path"]]
    wins = [["Kind", "Dataset", "Rank (1 wins)", "Rules path"]]
    for code, entry in datasets:
        contract = entry.get("contract") or {}
        adapter = contract.get("adapter") or {}
        at = f"mdm.{code}.contract"
        about += [
            [f"{code}: provider", _text(contract.get("provider"))],
            [f"{code}: kind", _text(adapter.get("kind") or (adapter.get("kind_values") and "by category: "
                                     + _text(adapter["kind_values"])) or (adapter.get("classification") or {}).get("kind"))],
            [f"{code}: one record is", _text(contract.get("record_key"))],
            [f"{code}: one publication is", _text(contract.get("publication_key"))],
            [f"{code}: a file covers", _text(contract.get("completeness") or contract.get("semantics"))],
            [f"{code}: set aside without stopping the run", _text(contract.get("nonblocking_deferred_reasons"))],
        ]
        for field, spec in (adapter.get("fields") or {}).items():
            for part, path in _paths(spec):
                fields.append([code, field, part, path, "MDM field", f"{at}.adapter.fields.{field}"])
        for field, spec in (adapter.get("matching") or {}).items():
            for part, path in _paths(spec):
                fields.append([code, field, part, path, "Matching only", f"{at}.adapter.matching.{field}"])
        for relationship in adapter.get("relationships") or []:
            kinds = relationship.get("type") or _text(list((relationship.get("type_values") or {}).values()))
            fields.append([code, f"relationship {kinds}", "to", _text(relationship.get("target_key")),
                           f"Link to {relationship.get('target_source')}", f"{at}.adapter.relationships"])
        formats = adapter.get("identifier_formats") or {}
        for ident, path in (adapter.get("identifiers") or {}).items():
            identifiers.append([code, ident, path, _text(formats.get(ident)),
                                "Yes" if ident in JOINING else "No: lookup only",
                                "Only with a new dataset code", f"{at}.adapter.identifiers.{ident}"])
        block = contract.get("quality") or {}
        for i, fix in enumerate(block.get("fixes") or []):
            args = fix.get("args") or {}
            quality.append([code, fix["id"], "Fix", _words(QUALITY_WORDS, fix["fix"], args),
                            _text(args.get("field") or args.get("name")), "",
                            f"quality.{code}.fixes[{i}]"])
        for i, check in enumerate(block.get("checks") or []):
            args = check.get("args") or {}
            quality.append([code, check["id"], "Check", _words(QUALITY_WORDS, check["test"], args),
                            check["value"], ON_FAIL_WORDS.get(check["on_fail"], check["on_fail"]),
                            f"quality.{code}.checks[{i}]"])
            if check["on_fail"] == "exception" and check["test"] == "present@1":
                critical.append([code, check["value"].removeprefix("fields."), ON_FAIL_WORDS["exception"],
                                 f"quality.{code}.checks[{i}]"])
        for kind, rules in (policy.get("kinds") or {}).items():
            ranked = (rules.get("defaults") or {}).get("sources") or []
            if code in ranked:
                wins.append([kind, code, str(ranked.index(code) + 1), f"merge/kinds/{kind}.yaml defaults.sources"])
    return {"Source": about, "Fields": fields, "Identifiers": identifiers,
            "Critical data elements": critical, "Data quality": quality, "Who wins": wins}


# --- a kind's workbook --------------------------------------------------------------

def _primitive(item: dict) -> str:
    return f"{item.get('primitive')}({_text(item.get('args') or {})})"


def _kind_sheets(kind: str, rules: dict) -> dict[str, list[list[str]]]:
    preferred = [["Rank (1 wins)", "Dataset", "Rules path"]]
    for i, code in enumerate((rules.get("defaults") or {}).get("sources") or []):
        preferred.append([str(i + 1), code, f"merge/kinds/{kind}.yaml defaults.sources[{i}]"])
    classification = [["Rule", "Version", "Dataset", "Step", "Verdict", "Probably", "When", "Rules path"]]
    matching = [["Order", "Rule", "Version", "Links", "To", "Every one of these must hold", "If nothing matches",
                 "Rules path"]]
    order = 0
    for i, rule in enumerate(rules.get("rules") or []):
        at = f"merge/kinds/{kind}.yaml rules[{i}]"
        if rule.get("family") == "classification":
            for j, step in enumerate(rule.get("steps") or []):
                when = "otherwise" if step.get("otherwise") else "; ".join(_primitive(w) for w in step.get("when") or [])
                classification.append([rule["rule_id"], rule.get("version", ""), rule.get("source", ""),
                                       _text(step.get("step")), _text(step.get("verdict")),
                                       _text(step.get("probable_kind")), when, f"{at}.steps[{j}]"])
        else:
            order += 1
            matching.append([str(order), rule["rule_id"], rule.get("version", ""), rule.get("source", ""),
                             rule.get("holder_source", ""), "; ".join(_primitive(w) for w in rule.get("when") or []),
                             _text(rule.get("on_no_match")), at])
    about = [["What", "Value"], ["How to change this", HOW_TO_CHANGE], ["Kind", kind],
             ["Rules version", _text(rules.get("version"))],
             ["Proof bar", _text(rules.get("bars"))]]
    return {"Kind": about, "Preferred sources": preferred, "Classification": classification,
            "Matching rules": matching}


# --- the stewards' notes -------------------------------------------------------------

def comment_notes(*paths: Path) -> list[list[str]]:
    """The decisions written as comments in rules files, one note per comment
    block, naming the key it sits above. A new workbook starts its Notes
    sheet with them: a rules export drops comments, a workbook keeps them."""
    notes = []
    for path in paths:
        if not path.exists():
            continue
        block: list[str] = []
        for line in path.read_text(encoding="utf-8").splitlines():
            text = line.strip()
            if text.startswith("#"):
                block.append(text.lstrip("#").strip())
            elif block and text:
                about = text.split(":", 1)[0].lstrip("- ").strip()
                notes.append([" ".join(b for b in block if b), f"{path.name}: {about}",
                              "rules file comment", ""])
                block = []
    return notes


# --- which workbooks there are ------------------------------------------------------

def documents(root: Path | None = None) -> dict[Path, tuple[dict[str, list[list[str]]], list[Path]]]:
    """Every Mapping Document the rules under `root` generate, by path, with
    the rules files it is generated from."""
    root = root or files.ROOT
    policy = files.policy(root)
    found: dict[Path, tuple] = {}
    for path in sorted((root / "sources").glob("*/source.yaml")):
        body = files.load_source(path)
        if body.get("mdm"):  # a source that feeds no MDM kind maps nothing
            found[path.parent / "MAPPING.xlsx"] = (
                _source_sheets(path.parent.name, body, policy), [path, path.parent / "quality.yaml"])
    for kind, rules in sorted((policy.get("kinds") or {}).items()):
        yaml = root / "merge" / "kinds" / f"{kind}.yaml"
        found[yaml.with_suffix(".xlsx")] = (_kind_sheets(kind, rules), [yaml])
    return found


# --- the workbook file ----------------------------------------------------------------

def read(path: Path) -> dict[str, list[list[str]]]:
    """A workbook's sheets as rows of text."""
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return {
            sheet.title: [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)
                          if any(v not in (None, "") for v in row)]
            for sheet in book.worksheets
        }
    finally:
        book.close()


def write(path: Path, sheets: dict[str, list[list[str]]], sources: list[Path] = ()) -> None:
    """Write a workbook, keeping the Notes sheet an existing one has; a new
    one starts its notes from the comments in its rules files."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    notes = read(path).get(NOTES) if path.exists() else [NOTES_HEADER, *comment_notes(*sources)]
    book = Workbook()
    book.remove(book.active)
    for title, rows in {**sheets, NOTES: notes or [NOTES_HEADER]}.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            width = max(len(str(c.value or "")) for c in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max(width, 10), 60) + 2
            for c in column:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def differences(path: Path, sheets: dict[str, list[list[str]]]) -> list[str]:
    """What a workbook says that its rules do not, cell by cell, in plain text.
    The Notes sheet is the stewards' own and never differs."""
    if not path.exists():
        return [f"{path}: missing; run `rules mapdoc write`"]
    found = read(path)
    out = []
    for title in sorted(set(sheets) | set(found) - {NOTES}):
        if title == NOTES:
            continue
        want, have = sheets.get(title), found.get(title)
        if want is None or have is None:
            out.append(f"{title}: sheet {'added' if want is None else 'removed'} in the workbook")
            continue
        header = want[0] if want else []
        for i in range(max(len(want), len(have))):
            a = want[i] if i < len(want) else None
            b = have[i] if i < len(have) else None
            if a is None:
                out.append(f"{title} row {i + 1}: added in the workbook: {' | '.join(b)}")
            elif b is None:
                out.append(f"{title} row {i + 1}: removed in the workbook: {' | '.join(a)}")
            else:
                for j in range(max(len(a), len(b))):
                    x = a[j] if j < len(a) else ""
                    y = b[j] if j < len(b) else ""
                    if x != y:
                        column = header[j] if j < len(header) else f"column {j + 1}"
                        out.append(f"{title} row {i + 1} ({a[0] if a else ''}), {column}: "
                                   f"rules say {json.dumps(x)}, workbook says {json.dumps(y)}")
    return out
