"""Rules skill ticket 11: the Mapping Document, a spreadsheet held equal to the rules."""

from __future__ import annotations

import hashlib
import shutil
from argparse import Namespace

import pytest

from edgar_warehouse.rules import files, mapdoc


def test_every_committed_mapping_document_equals_its_rules():
    """A workbook that differs from the rules it is generated from fails here:
    a steward's change must become rules, or be regenerated away, before it
    merges. Fix with `edgar-warehouse rules mapdoc write`."""
    for path, (_, sheets, _) in mapdoc.documents().items():
        assert mapdoc.differences(path, sheets) == [], path


def test_only_sources_that_feed_mdm_have_one():
    assert {name for name, _, _ in mapdoc.documents().values()} == {"gleif", "sec.submissions.company", "company"}


@pytest.fixture
def rules_copy(tmp_path):
    root = tmp_path / "rules"
    shutil.copytree(files.ROOT, root)
    return root


def _write_all(root):
    for path, (_, sheets, sources) in mapdoc.documents(root).items():
        mapdoc.write(path, sheets, sources)
    return mapdoc.documents(root)


def _edit(path, change):
    from openpyxl import load_workbook

    book = load_workbook(path)
    change(book)
    book.save(path)


def test_a_steward_change_is_reported_cell_by_cell(rules_copy):
    found = _write_all(rules_copy)
    path = rules_copy / "sources" / "gleif" / "MAPPING.xlsx"

    def change(book):
        book["Critical data elements"].append(["gleif.level1.v1", "address", "The value is filled", "Exception", ""])
        book["Identifiers"]["B2"] = "lei_code"

    _edit(path, change)
    lines = mapdoc.differences(path, found[path][1])
    assert lines == [
        "Critical data elements row 3: added in the workbook: gleif.level1.v1 | address | The value is filled | Exception",
        'Identifiers row 2 (gleif.level1.v1), Identifier: rules say "lei", workbook says "lei_code"',
    ]


def test_one_row_inserted_is_one_line_not_every_row_below_it(rules_copy):
    found = _write_all(rules_copy)
    path = rules_copy / "merge" / "kinds" / "company.xlsx"
    _edit(path, lambda book: book["Classification"].insert_rows(3))
    _edit(path, lambda book: [book["Classification"].cell(3, 1, "new-rule")])
    assert mapdoc.differences(path, found[path][1]) == [
        "Classification row 3: added in the workbook: new-rule"
    ]


def test_notes_are_kept_when_the_workbook_is_regenerated(rules_copy):
    found = _write_all(rules_copy)
    path = rules_copy / "merge" / "kinds" / "company.xlsx"
    _edit(path, lambda book: book[mapdoc.NOTES].append(
        ["Rank SEC first for names", "Preferred sources", "a steward", "2026-09-28"]))
    _, sheets, sources = found[path]
    mapdoc.write(path, sheets, sources)
    assert ["Rank SEC first for names", "Preferred sources", "a steward", "2026-09-28"] in (
        mapdoc.read(path)[mapdoc.NOTES])
    assert mapdoc.differences(path, sheets) == []


def test_a_new_workbook_keeps_the_decisions_the_rules_comments_hold(rules_copy):
    found = mapdoc.documents(rules_copy)
    path = rules_copy / "sources" / "sec.submissions.company" / "MAPPING.xlsx"
    path.unlink(missing_ok=True)
    _, sheets, sources = found[path]
    mapdoc.write(path, sheets, sources)
    notes = mapdoc.read(path)[mapdoc.NOTES]
    assert notes[0] == mapdoc.NOTES_HEADER
    assert any(row[1] == "source.yaml: identifiers" and "EIN" in row[0] for row in notes)


def test_regenerating_unchanged_rules_writes_the_same_file(rules_copy):
    path, (_, sheets, sources) = next(iter(mapdoc.documents(rules_copy).items()))
    mapdoc.write(path, sheets, sources)
    first = hashlib.sha256(path.read_bytes()).hexdigest()
    mapdoc.write(path, sheets, sources)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == first


def test_text_that_looks_like_a_formula_stays_text(tmp_path):
    path = tmp_path / "x.xlsx"
    mapdoc.write(path, {"Sheet": [["What"], ["=1+1"]]})
    assert mapdoc.read(path)["Sheet"] == [["What"], ["=1+1"]]


def test_the_sheets_say_what_the_rules_say():
    found = {name: sheets for name, sheets, _ in mapdoc.documents().values()}
    sec = found["sec.submissions.company"]
    assert ["sec.submissions.company.v1", "cik", "cik", "sec_cik", "Yes", "Only with a new dataset code",
            "mdm.sec.submissions.company.v1.contract.adapter.identifiers.cik"] in sec["Identifiers"]
    assert [row[1] for row in sec["Critical data elements"][1:]] == ["name"]
    assert ["company", "name", "sec.submissions.company.v1, gleif.level1.v1", "sec.submissions.company.v1",
            "Kind default", "merge/kinds/company.yaml defaults.sources"] in sec["Who wins"]
    kind = found["company"]
    assert [row[1] for row in kind["Matching rules"][1:]] == [
        "company-cik", "sec-gleif-name-jurisdiction", "sec-gleif-name-postal",
        *(f"sec-gleif-cascade-p{n}" for n in range(1, 8))]
    assert "Exactly one Company holds that cik" in kind["Matching rules"][1][5]
    assert "the name alone" in kind["Matching rules"][10][5] and "the country, street" in kind["Matching rules"][4][5]
    assert "The SEC company is not already linked to another LEI" in kind["Matching rules"][2][5]
    assert ["jurisdiction", "gleif.level1.v1", "gleif.level1.v1", "Kind default",
            "merge/kinds/company.yaml defaults.sources"] in kind["Who wins each field"]
    assert len(kind["Classification"]) == 15  # the 14 steps of sec-company-candidate


def test_every_quality_test_and_fix_has_plain_words():
    from edgar_warehouse.mdm.clean.quality import CHECKS, FIXES

    assert set(mapdoc.QUALITY_WORDS) == set(CHECKS) | set(FIXES)


def test_every_condition_the_rules_use_has_plain_words():
    used = set()
    for rules in files.policy()["kinds"].values():
        for rule in rules["rules"]:
            used |= {w["primitive"] for w in rule.get("when") or []}
            used |= {w["primitive"] for s in rule.get("steps") or [] for w in s.get("when") or []}
    assert used <= set(mapdoc.CONDITION_WORDS)


def test_a_long_list_shows_as_a_count_and_every_cell_is_text():
    assert mapdoc._text(list(range(13))) == "13 values (see the rules file)"
    assert mapdoc._text(["A", "B"]) == "A, B"
    assert mapdoc._text(3) == "3"


def test_only_selects_by_source_or_kind_name(capsys):
    from edgar_warehouse.rules.cli import _mapdoc

    args = Namespace(root=str(files.ROOT), action="check")
    assert _mapdoc(Namespace(**vars(args), only="kinds")) == 2
    assert _mapdoc(Namespace(**vars(args), only="gleif")) == 0


def test_a_field_with_its_own_rule_shows_its_own_winner(rules_copy):
    """Rules skill ticket 12: the merge engine takes a rule per field
    (`fields.<name>`), so the workbook shows it, and a steward's change to
    it is reported."""
    kind_file = rules_copy / "merge" / "kinds" / "company.yaml"
    body = files.load(kind_file)
    body["fields"] = {"address": {"sources": ["gleif.level1.v1", "sec.submissions.company.v1"]}}
    kind_file.write_text(files.dumps(body), encoding="utf-8")
    found = _write_all(rules_copy)
    by_name = {name: sheets for name, sheets, _ in found.values()}
    assert ["address", "gleif.level1.v1, sec.submissions.company.v1", "gleif.level1.v1", "Its own rule",
            "merge/kinds/company.yaml fields.address"] in by_name["company"]["Who wins each field"]
    assert ["name", "sec.submissions.company.v1, gleif.level1.v1", "sec.submissions.company.v1", "Kind default",
            "merge/kinds/company.yaml defaults.sources"] in by_name["company"]["Who wins each field"]
    assert ["company", "address", "gleif.level1.v1, sec.submissions.company.v1", "gleif.level1.v1",
            "Its own rule", "merge/kinds/company.yaml fields.address"] in by_name["gleif"]["Who wins"]
    path = rules_copy / "merge" / "kinds" / "company.xlsx"
    rows = found[path][1]["Who wins each field"]
    row = next(i for i, r in enumerate(rows) if r[0] == "address") + 1
    _edit(path, lambda book: book["Who wins each field"].cell(row, 3, "sec.submissions.company.v1"))
    assert mapdoc.differences(path, found[path][1]) == [
        f'Who wins each field row {row} (address), Winner: rules say "gleif.level1.v1", '
        'workbook says "sec.submissions.company.v1"'
    ]


def test_a_field_rule_without_its_own_order_says_the_order_is_the_default(rules_copy):
    kind_file = rules_copy / "merge" / "kinds" / "company.yaml"
    body = files.load(kind_file)
    body["fields"] = {"name": {"max_age_days": 400}}
    kind_file.write_text(files.dumps(body), encoding="utf-8")
    sheets = {name: s for name, s, _ in mapdoc.documents(rules_copy).values()}["company"]
    assert ["name", "sec.submissions.company.v1, gleif.level1.v1", "sec.submissions.company.v1",
            "Its own rule (max_age_days: 400), order from the kind default",
            "merge/kinds/company.yaml fields.name; merge/kinds/company.yaml defaults.sources"] in (
        sheets["Who wins each field"])
